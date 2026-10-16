# -*- coding: utf-8 -*-
"""Preliminares: portada, ficha de identificación, control de versiones y glosario."""
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH


def escribir(M, P):
    doc = M.doc

    # ================================================================ PORTADA
    for _ in range(2):
        M.p("", espacio_despues=0)

    par = M.p("SERVICIO NACIONAL DE APRENDIZAJE – SENA", negrita=True, centrado=True, espacio_despues=2)
    for r in par.runs:
        r.font.size = Pt(14)
    M.p("República de Colombia", centrado=True, espacio_despues=18)
    for r in M.doc.paragraphs[-1].runs:
        r.font.size = Pt(10)

    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.add_run().add_picture(P["logo"], width=Cm(6.2))

    M.p("", espacio_despues=10)
    p = M.p("MANUAL TÉCNICO", negrita=True, centrado=True, espacio_despues=4)
    for r in p.runs:
        r.font.size = Pt(16)
    p = M.p("ProyecTwin", negrita=True, centrado=True, espacio_despues=4)
    for r in p.runs:
        r.font.size = Pt(16)
    M.p("Plataforma de análisis de similitud para proyectos formativos de aprendices SENA",
        centrado=True, cursiva=True, espacio_despues=26)

    filas_portada = [
        ("Versión del sistema:", "MVP web (React + Laravel) y aplicación Android 1.0 — el repositorio no declara un número de versión semántico único (ver § 2.1)"),
        ("Equipo de desarrollo:", "Autores registrados en el historial de Git del repositorio: «Oscar» y «YDanB-312». La nómina completa, la ficha y los datos institucionales no constan en los archivos del proyecto (ver § 1.9 y Anexo F)"),
        ("Ciudad y año:", "Popayán, Cauca — Colombia · Octubre de 2026"),
    ]
    for etiqueta, valor in filas_portada:
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = par.add_run(etiqueta + " ")
        r.bold = True
        r.font.size = Pt(10.5)
        r2 = par.add_run(valor)
        r2.font.size = Pt(10.5)
        par.paragraph_format.space_after = Pt(3)

    M.p("", espacio_despues=22)
    M.p("Documento generado a partir de la auditoría técnica del repositorio realizada el 16 de octubre de 2026. "
        "Las afirmaciones técnicas se sustentan en el código fuente, los archivos de configuración y las verificaciones ejecutadas que se relacionan en los anexos.",
        centrado=True, cursiva=True, espacio_despues=0)

    # ================================================================ FICHA
    sec2 = M.salto_seccion()
    M.encabezado_y_pie(sec2, "Manual Técnico — ProyecTwin", numeracion="roman", inicio=1)
    M.h1("Ficha de identificación del documento")

    M.tabla(
        "Ficha técnica del documento",
        ["Campo", "Valor"],
        [
            ("Nombre del proyecto", "ProyecTwin (MVP) — plataforma de análisis de similitud para proyectos formativos"),
            ("Nombre del documento", "Manual Técnico de ProyecTwin"),
            ("Identificador documental", "PT-MT-001"),
            ("Versión del manual", "1.0"),
            ("Fecha de elaboración", "16 de octubre de 2026"),
            ("Estado del documento", "Completo para revisión del equipo; generado de forma automatizada y pendiente de validación humana final"),
            ("Ámbito técnico", "Backend Laravel 10, frontend React 19, aplicación Android (Kotlin), base de datos MySQL, motor de similitud, seguridad, pruebas y despliegue del MVP"),
            ("Equipo de desarrollo", "«Oscar» y «YDanB-312» (autores de los 46 commits del repositorio). Nómina completa, roles individuales y ficha de formación: no constan en los archivos del proyecto"),
            ("Programa de formación", "No confirmado en el repositorio. El sistema gestiona programas SENA (p. ej., ADSO) como datos de su catálogo, lo que no acredita el programa del equipo desarrollador"),
            ("Responsable de elaboración", "Auditoría y generación documental automatizada (asistida por OpenCode) sobre el repositorio real; revisión y correcciones a cargo del equipo del proyecto"),
            ("Repositorio y rama", "https://github.com/YDanB-312/ProyecTwin-MVP — rama master"),
            ("Revisión auditada", "46 commits; último commit registrado: 3b15f3f (13/10/2026). Al momento de la auditoría existían cambios sin confirmar, documentados en el Anexo F"),
            ("Entorno de verificación", "Windows, PHP 8.2.12 (XAMPP), Composer 2.10.2, Node 26.3.1, MariaDB 10.4.32, Microsoft Word disponible para validación del DOCX"),
        ],
        fuente="Elaboración propia (auditoría del repositorio ProyecTwin-MVP, 16/10/2026).",
        anchos=[4.6, 11.4],
    )

    M.h2("Nota de alcance documental (discrepancia entre fuentes)")
    M.p("El material de referencia entregado junto con la solicitud de este manual incluía el documento "
        "«M.U- PopayanAllTour -SdeI.pdf», un manual de usuario de **PopayanAllTour**, un proyecto distinto "
        "(aplicación turística de Popayán) que no guarda relación con el código de este repositorio y cuyo código "
        "fuente no existe en el entorno de trabajo. Siguiendo la regla de prioridad de fuentes —el código y los "
        "archivos de configuración determinan lo que está implementado—, este manual documenta exclusivamente el "
        "sistema **ProyecTwin-MVP** y registra la discrepancia en el Anexo F. Ninguna afirmación técnica de este "
        "documento proviene del PDF de PopayanAllTour.")
    M.p("Este documento es una **verificación de ingeniería basada en evidencia** (código, configuración y "
        "ejecución local). No constituye una auditoría de certificación, una prueba de penetración ni una "
        "declaración de conformidad normativa.")

    # ================================================== CONTROL DE VERSIONES
    M.doc.add_page_break()
    M.h1("Control de versiones del documento")
    M.p("El siguiente registro documenta la elaboración de este manual. Las filas finales quedan reservadas "
        "para que el equipo del proyecto registre sus propias revisiones.")
    M.tabla(
        "Control de versiones del manual técnico",
        ["Versión", "Fecha", "Descripción del cambio", "Responsable"],
        [
            ("0.1", "16/10/2026", "Auditoría inicial del repositorio, inventario de evidencias y definición de la estructura documental", "Generación documental asistida (OpenCode)"),
            ("0.5", "16/10/2026", "Ejecución del entorno local (backend, frontend, base de datos), pruebas verificadas y capturas reales de la interfaz", "Generación documental asistida (OpenCode)"),
            ("0.9", "16/10/2026", "Redacción de los ocho capítulos, figuras, tablas, referencias y anexos", "Generación documental asistida (OpenCode)"),
            ("1.0", "16/10/2026", "Versión completa entregada para revisión: campos de índice y listas activados, control de calidad documental ejecutado", "Generación documental asistida (OpenCode)"),
            ("—", "—", "(Reservado: correcciones y validación del equipo antes de la entrega institucional)", "Equipo ProyecTwin"),
        ],
        fuente="Elaboración propia.",
        anchos=[1.8, 2.2, 9.4, 2.6],
    )

    # =========================================== TABLA DE CONTENIDO Y LISTAS
    M.doc.add_page_break()
    M.h1("Tabla de contenido")
    M.p("La tabla siguiente es un campo automático de Word. Si aparece sin diligenciar, actualícela con "
        "«Referencias → Actualizar tabla» (o `F9` sobre el campo).")
    M.campo_toc()

    M.doc.add_page_break()
    M.h1("Listas de figuras y tablas")
    M.p("Las listas se generan a partir de los campos de numeración de cada leyenda (actualizar si aparecen "
        "vacías).")
    M.h2("Índice de figuras")
    M.campo_lista("Figura")
    M.h2("Índice de tablas")
    M.campo_lista("Tabla")

    # ================================================================ GLOSARIO
    M.doc.add_page_break()
    M.h1("Glosario de términos y siglas")
    M.p("Definiciones de los términos y siglas empleados en este manual, alineadas con las tecnologías que el "
        "sistema utiliza realmente.")

    glosario = [
        ("ADSO", "Análisis y Desarrollo de Software: programa de formación del SENA presente en el catálogo del sistema (no acreditado como programa del equipo)."),
        ("Aprendiz / Instructor / Administrador", "Los tres roles funcionales del sistema, definidos en la columna `rol` de `general_users` y exigidos por el middleware `rol`."),
        ("API REST", "Interfaz HTTP con recursos y verbos estándar. ProyecTwin expone su API bajo el prefijo `/v1` mediante Laravel."),
        ("Archivado de similitud", "Transición de un par detectado a `vigente = false` para conservarlo como evidencia histórica cuando el contenido cambia o la propuesta se rechaza."),
        ("Bearer token", "Token de acceso enviado en la cabecera `Authorization`. Sanctum lo emite para clientes que no usan cookie (app Android y pruebas automatizadas)."),
        ("BM25", "Función de ranking de recuperación de información usada por el motor como primera señal (peso 0,40)."),
        ("Corpus (del motor)", "Conjunto de propuestas aprobadas del mismo programa, dentro de la ventana temporal configurada, contra el que se compara cada propuesta enviada."),
        ("CSRF", "Cross-Site Request Forgery. Laravel lo mitiga con cookie `XSRF-TOKEN` y la cabecera `X-XSRF-TOKEN` en cada escritura de la SPA."),
        ("Cookie httpOnly", "Cookie de sesión no legible desde JavaScript; es el mecanismo de sesión principal de la SPA web."),
        ("Eloquent", "ORM de Laravel: cada tabla tiene un modelo (`app/Models`) y relaciones declaradas."),
        ("Ficha (class group)", "Grupo de formación SENA. Tiene código de unión, instructor a cargo, programa, y estados `activo`, `finalizado` y `anulada`."),
        ("Huella de envío", "Resumen `sha1` del contenido de una propuesta (`projects.huella_envio`) que impide reenviar sin cambios reales."),
        ("LSA", "Latent Semantic Analysis (análisis semántico latente). Señal temática del motor implementada en PHP puro (`app/Similarity/Lsa.php`)."),
        ("Middleware", "Componente que intercepta la petición antes del controlador. ProyecTwin usa `auth:sanctum`, `cuenta.activa`, `CambioContrasenaObligatorio` y `rol`."),
        ("Migración", "Archivo versionado que crea o modifica la estructura de la base de datos (`database/migrations`)."),
        ("Motor de similitud", "Conjunto de clases PHP (`SimilitudService`, `Recomputador`, `Lsa`, `Pasajes`) que compara propuestas y produce pares con porcentaje, señales y pasajes."),
        ("Par de similitud", "Fila de `similarities` con dos propuestas (`id_proyecto_1 < id_proyecto_2`), porcentaje 0–100, detalles y bandera `vigente`."),
        ("Pasajes", "Fragmentos textuales coincidentes entre dos propuestas, estilo herramienta antiplagio (hasta 5 pares por comparación)."),
        ("Pivote equipo", "Tabla `apprentice_projects`: relación N:M entre aprendices y propuestas; define quién participa en cada proyecto."),
        ("PWA", "Progressive Web App: el frontend incluye `manifest.webmanifest` para instalación básica."),
        ("Regla «Classroom»", "Regla de negocio implementada en `Project::puedeEscribir`: solo se participa (editar, comentar, gestionar equipo) dentro de la ficha activa; fuera de ella, solo lectura."),
        ("Sanctum", "Paquete de Laravel para autenticar SPA (cookie) y clientes móviles (token) contra la misma API."),
        ("Scoping", "Filtrado de datos por alcance del usuario (ficha, autoría, asignación) codificado en los `scope*` de los modelos (p. ej. `Project::scopeParaUsuario`)."),
        ("Seeder", "Clase que carga datos de demostración (`DatabaseSeeder`) o de volumen (`VolumenSeeder`)."),
        ("SPA", "Single Page Application: el frontend React se carga una vez y consume la API por `fetch`."),
        ("Umbral", "Porcentaje mínimo a partir del cual un par se considera similar; se configura en `motor_configs` (por defecto 0,30 para el seed)."),
        ("Ventana temporal", "Cantidad de meses hacia atrás que delimita el corpus del motor (`motor_configs.meses`, por defecto 12)."),
        ("Vitest / Playwright / PHPUnit", "Herramientas de prueba usadas por el proyecto: unitarias de frontend, extremo a extremo de navegador y pruebas de backend, respectivamente."),
        ("Word / DOCX", "Formato de este manual; los índices y listas usan campos automáticos que Word actualiza al abrir el documento."),
    ]
    M.tabla(
        "Glosario técnico",
        ["Término", "Definición"],
        glosario,
        fuente="Elaboración propia con base en el código auditado.",
        anchos=[4.2, 11.8],
        tamano=9,
    )
