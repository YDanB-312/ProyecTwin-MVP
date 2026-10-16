# -*- coding: utf-8 -*-
"""
manual_lib.py — Infraestructura de generación del Manual Técnico de ProyecTwin.

Construye el documento Word con:
  • Estilos nativos (Título, Título 1-3, Normal, Leyenda) en Arial según la
    especificación institucional (Arial 11; H1 13 negrita; H2 12 negrita).
  • Secciones con numeración preliminar romana y numeración árabe en el cuerpo.
  • Campos automáticos reales de Word: TOC, listas de figuras y tablas
    (campos SEQ en las leyendas) y PAGE en el pie.
  • Tablas con encabezado repetible, sombreado sobrio y tipografía Arial 9.
  • Activa la actualización de campos al abrir (w:updateFields).

Autor: generación documental de la auditoría (octubre de 2026).
Uso exclusivo para construir el manual; no modifica el código de la aplicación.
"""
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu

AZUL_INST = RGBColor(0x1F, 0x2A, 0x33)      # texto de títulos
VERDE_SENA = "2A7C00"                        # encabezados de tabla
VERDE_CLARO = "F3F9EF"                       # filas alternas
GRIS_SUAVE = "F2F4F5"                        # fondo de código
GRIS_BORDE = "C8CED4"
TEXTO = RGBColor(0x22, 0x2B, 0x33)


def _set(el, tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), v)
    el.append(e)
    return e


def _sombra_celda(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _sombra_parrafo(par, fill):
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def _borde_parrafo(par, lado="left", color=VERDE_SENA, sz=18):
    pPr = par._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    b = OxmlElement(f"w:{lado}")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(sz))
    b.set(qn("w:space"), "6")
    b.set(qn("w:color"), color)
    pbdr.append(b)
    pPr.append(pbdr)


def forzar_fuente(estilo, fuente):
    """Fija la fuente de un estilo eliminando las referencias al tema (que en
    OOXML tienen prioridad sobre los atributos ascii/hAnsi y provocarían que
    Word ignore la fuente indicada)."""
    estilo.font.name = fuente  # asegura la existencia de w:rFonts
    rPr = estilo.element.find(qn("w:rPr"))
    if rPr is None:
        return
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        return
    for attr in ("asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme"):
        clave = qn("w:" + attr)
        if rFonts.get(clave) is not None:
            del rFonts.attrib[clave]
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        rFonts.set(qn("w:" + attr), fuente)


def _campo(par, instruccion, placeholder=""):
    """Inserta un campo de Word (p. ej. TOC, PAGE, SEQ) en el párrafo."""
    r1 = par.add_run()
    _set(r1._r, "w:fldChar", **{"w:fldCharType": "begin"})
    r2 = par.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instruccion} "
    r2._r.append(it)
    r3 = par.add_run()
    _set(r3._r, "w:fldChar", **{"w:fldCharType": "separate"})
    r4 = par.add_run(placeholder)
    r5 = par.add_run()
    _set(r5._r, "w:fldChar", **{"w:fldCharType": "end"})
    return [r1, r2, r3, r4, r5]


class Manual:
    def __init__(self, ruta_raiz):
        self.ruta = ruta_raiz
        self.doc = Document()
        self.n_figuras = 0
        self.n_tablas = 0
        self.n_codigos = 0
        self._configurar_pagina(self.doc.sections[0], primera=True)
        self._crear_estilos()
        self._activar_update_fields()

    # ------------------------------------------------------------- estilos
    def _configurar_pagina(self, sec, primera=False):
        sec.page_width = Cm(21.59)   # Carta 8.5 in
        sec.page_height = Cm(27.94)  # Carta 11 in
        for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, lado, Cm(2.5))
        sec.header_distance = Cm(1.2)
        sec.footer_distance = Cm(1.2)

    def _crear_estilos(self):
        est = self.doc.styles

        def get_or_create(nombre, tipo=WD_STYLE_TYPE.PARAGRAPH, base=None):
            try:
                return est[nombre]
            except KeyError:
                s = est.add_style(nombre, tipo)
                if base is not None:
                    s.base_style = est[base]
                return s

        # Normal — cuerpo del texto
        n = est["Normal"]
        n.font.size = Pt(11)
        n.font.color.rgb = TEXTO
        n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        n.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        n.paragraph_format.space_after = Pt(6)
        forzar_fuente(n, "Arial")

        # Estilos de título (nativos, sobrescritos)
        h1 = est["Heading 1"]
        h1.font.size = Pt(13)
        h1.font.bold = True
        h1.font.color.rgb = AZUL_INST
        h1.paragraph_format.space_before = Pt(18)
        h1.paragraph_format.space_after = Pt(8)
        h1.paragraph_format.keep_with_next = True
        h1.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

        h2 = est["Heading 2"]
        h2.font.size = Pt(12)
        h2.font.bold = True
        h2.font.color.rgb = AZUL_INST
        h2.paragraph_format.space_before = Pt(12)
        h2.paragraph_format.space_after = Pt(6)
        h2.paragraph_format.keep_with_next = True
        h2.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

        h3 = est["Heading 3"]
        h3.font.size = Pt(11)
        h3.font.bold = True
        h3.font.italic = False
        h3.font.color.rgb = AZUL_INST
        h3.paragraph_format.space_before = Pt(10)
        h3.paragraph_format.space_after = Pt(4)
        h3.paragraph_format.keep_with_next = True
        h3.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

        # Forzar Arial (quitando referencias al tema, que tienen prioridad)
        for estilo in (h1, h2, h3):
            forzar_fuente(estilo, "Arial")

        t = est["Title"]
        t.font.size = Pt(16)
        t.font.bold = True
        t.font.color.rgb = AZUL_INST
        forzar_fuente(t, "Arial")

        # Estilos de las entradas del índice (TOC 1-3), si existen
        for nombre_toc in ("TOC 1", "TOC 2", "TOC 3"):
            try:
                forzar_fuente(est[nombre_toc], "Arial")
            except KeyError:
                pass

        # Leyenda de figuras/tablas
        ley = get_or_create("LeyendaManual", base="Caption")
        ley.font.size = Pt(9)
        ley.font.italic = True
        ley.font.bold = False
        ley.font.color.rgb = RGBColor(0x44, 0x4F, 0x5A)
        ley.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ley.paragraph_format.space_before = Pt(4)
        ley.paragraph_format.space_after = Pt(4)
        ley.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        forzar_fuente(ley, "Arial")

        # Fuente (bajo figura/tabla)
        fuente = get_or_create("FuenteManual", base="Caption")
        fuente.font.size = Pt(8)
        fuente.font.italic = True
        fuente.font.color.rgb = RGBColor(0x66, 0x6F, 0x77)
        fuente.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fuente.paragraph_format.space_before = Pt(2)
        fuente.paragraph_format.space_after = Pt(10)
        fuente.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        forzar_fuente(fuente, "Arial")

        # Texto de tablas
        tt = get_or_create("TablaTexto")
        tt.font.size = Pt(9)
        tt.font.color.rgb = TEXTO
        tt.paragraph_format.space_after = Pt(0)
        tt.paragraph_format.space_before = Pt(0)
        tt.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        forzar_fuente(tt, "Arial")

        ttb = get_or_create("TablaTitulo")
        ttb.font.size = Pt(9)
        ttb.font.bold = True
        ttb.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        ttb.paragraph_format.space_after = Pt(0)
        ttb.paragraph_format.space_before = Pt(0)
        ttb.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        forzar_fuente(ttb, "Arial")

        # Código fuente
        c = get_or_create("CodigoManual")
        c.font.size = Pt(9)
        c.font.color.rgb = RGBColor(0x27, 0x31, 0x3C)
        c.paragraph_format.space_after = Pt(0)
        c.paragraph_format.space_before = Pt(0)
        c.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        c.paragraph_format.left_indent = Cm(0.5)
        c.paragraph_format.right_indent = Cm(0.2)
        forzar_fuente(c, "Consolas")

    def _activar_update_fields(self):
        settings = self.doc.settings.element
        uf = OxmlElement("w:updateFields")
        uf.set(qn("w:val"), "true")
        settings.append(uf)

    # ------------------------------------------------------------- secciones
    def salto_seccion(self, preliminar=False):
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._configurar_pagina(sec)
        return sec

    def encabezado_y_pie(self, sec, texto_encabezado, numeracion="decimal", inicio=1):
        sec.header.is_linked_to_previous = False
        sec.footer.is_linked_to_previous = False
        p = sec.header.paragraphs[0]
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(texto_encabezado)
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0x66, 0x6F, 0x77)
        _borde_parrafo(p, lado="bottom", color=GRIS_BORDE, sz=6)

        pf = sec.footer.paragraphs[0]
        pf.text = ""
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = pf.add_run("Página ")
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0x66, 0x6F, 0x77)
        _campo(pf, "PAGE", "1")
        for run in pf.runs:
            run.font.name = "Arial"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(0x66, 0x6F, 0x77)

        # Formato de numeración de la sección (si inicio es None, continúa la secuencia)
        sectPr = sec._sectPr
        pg = sectPr.find(qn("w:pgNumType"))
        if pg is None:
            pg = OxmlElement("w:pgNumType")
            sectPr.append(pg)
        pg.set(qn("w:fmt"), "lowerRoman" if numeracion == "roman" else "decimal")
        if inicio is not None:
            pg.set(qn("w:start"), str(inicio))

    # ------------------------------------------------------------- bloques
    def h1(self, texto, salto=False):
        if salto:
            self.doc.add_page_break()
        return self.doc.add_heading(texto, level=1)

    def h2(self, texto):
        return self.doc.add_heading(texto, level=2)

    def h3(self, texto):
        return self.doc.add_heading(texto, level=3)

    def p(self, texto, negrita=False, cursiva=False, centrado=False, espacio_despues=6):
        par = self.doc.add_paragraph()
        if centrado:
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(espacio_despues)
        self._runs_formato(par, texto, negrita, cursiva)
        return par

    def _runs_formato(self, par, texto, negrita=False, cursiva=False):
        """Soporta **negrita**, *cursiva* y `código en línea` dentro del texto."""
        import re
        partes = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", texto)
        for parte in partes:
            if not parte:
                continue
            es_codigo = False
            if parte.startswith("**") and parte.endswith("**"):
                r = par.add_run(parte[2:-2])
                r.bold = True
            elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
                r = par.add_run(parte[1:-1])
                r.italic = True
            elif parte.startswith("`") and parte.endswith("`") and len(parte) > 2:
                r = par.add_run(parte[1:-1])
                es_codigo = True
            else:
                r = par.add_run(parte)
            r.font.name = "Consolas" if es_codigo else "Arial"
            if negrita:
                r.bold = True
            if cursiva:
                r.italic = True

    def bullets(self, items, estilo="List Bullet"):
        for it in items:
            par = self.doc.add_paragraph(style=estilo)
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            self._runs_formato(par, it)
            for r in par.runs:
                r.font.name = "Arial"
                r.font.size = Pt(11)

    def numerada(self, items):
        self.bullets(items, estilo="List Number")

    def codigo(self, texto, archivo=None, ref=None):
        self.n_codigos += 1
        lineas = texto.strip("\n").split("\n")
        par = self.doc.add_paragraph(style="CodigoManual")
        for i, linea in enumerate(lineas):
            r = par.add_run(linea)
            r.font.name = "Consolas"
            if i < len(lineas) - 1:
                r.add_break()
        _sombra_parrafo(par, GRIS_SUAVE)
        _borde_parrafo(par, lado="left", color=VERDE_SENA, sz=12)
        if archivo:
            pie = self.doc.add_paragraph(style="FuenteManual")
            pie.alignment = WD_ALIGN_PARAGRAPH.LEFT
            etiqueta = ref or f"Fragmento {self.n_codigos}"
            pie.add_run(f"{etiqueta} — Archivo: {archivo}")
        return par

    def figura(self, archivo_png, titulo, fuente, ancho_cm=15.5, alt=None):
        self.n_figuras += 1
        numero = self.n_figuras
        cap = self.doc.add_paragraph(style="LeyendaManual")
        cap.paragraph_format.keep_with_next = True
        cap.add_run("Figura ")
        _campo(cap, "SEQ Figura \\* ARABIC", str(numero))
        cap.add_run(f". {titulo}")
        for r in cap.runs:
            r.font.name = "Arial"
            r.font.size = Pt(9)
            r.font.italic = True

        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(4)
        par.paragraph_format.space_after = Pt(2)
        run = par.add_run()
        run.add_picture(archivo_png, width=Cm(ancho_cm))
        # Texto alternativo (accesibilidad)
        try:
            inline = run._r.findall(qn("w:drawing"))[0][0][0]
            docPr = inline.find(qn("wp:docPr"))
            if docPr is not None:
                docPr.set("descr", alt or titulo)
                docPr.set("title", titulo)
        except Exception:
            pass

        fp = self.doc.add_paragraph(style="FuenteManual")
        fp.add_run(f"Fuente: {fuente}")
        return numero

    def tabla(self, titulo, encabezados, filas, fuente=None, anchos=None, tamano=9):
        self.n_tablas += 1
        numero = self.n_tablas
        cap = self.doc.add_paragraph(style="LeyendaManual")
        cap.paragraph_format.keep_with_next = True
        cap.add_run("Tabla ")
        _campo(cap, "SEQ Tabla \\* ARABIC", str(numero))
        cap.add_run(f". {titulo}")
        for r in cap.runs:
            r.font.name = "Arial"
            r.font.size = Pt(9)
            r.font.italic = True

        n_cols = len(encabezados)
        t = self.doc.add_table(rows=1, cols=n_cols)
        t.style = self.doc.styles["Table Grid"]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = True

        # Encabezado
        hdr = t.rows[0]
        trPr = hdr._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trPr.append(th)
        for j, texto in enumerate(encabezados):
            celda = hdr.cells[j]
            celda.text = ""
            par = celda.paragraphs[0]
            par.style = self.doc.styles["TablaTitulo"]
            par.add_run(texto)
            _sombra_celda(celda, VERDE_SENA)

        # Filas
        for i, fila in enumerate(filas):
            row = t.add_row()
            for j, valor in enumerate(fila):
                celda = row.cells[j]
                celda.text = ""
                par = celda.paragraphs[0]
                par.style = self.doc.styles["TablaTexto"]
                self._runs_formato(par, str(valor))
                for r in par.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(tamano)
                if i % 2 == 1:
                    _sombra_celda(celda, VERDE_CLARO)

        if anchos:
            for j, ancho in enumerate(anchos):
                for row in t.rows:
                    row.cells[j].width = Cm(ancho)

        if fuente:
            fp = self.doc.add_paragraph(style="FuenteManual")
            fp.add_run(f"Fuente: {fuente}")
        else:
            self.doc.add_paragraph(style="FuenteManual")
        return numero

    # ------------------------------------------------------------- campos
    def campo_toc(self):
        par = self.doc.add_paragraph()
        _campo(par, 'TOC \\o "1-3" \\h \\z \\u',
               "Tabla de contenido: actualice el campo (Referencias → Actualizar tabla) si no aparece diligenciada.")
        return par

    def campo_lista(self, tipo):
        """tipo: 'Figura' o 'Tabla' → lista automática por campos SEQ."""
        par = self.doc.add_paragraph()
        _campo(par, f'TOC \\h \\z \\c "{tipo}"',
               f"Lista de {tipo.lower()}s: actualice el campo si no aparece diligenciada.")
        return par

    def guardar(self, ruta_salida):
        self.doc.save(ruta_salida)
        return ruta_salida
